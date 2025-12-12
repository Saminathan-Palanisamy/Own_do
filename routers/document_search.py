# routers/document_search.py
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException, status
from sqlalchemy.orm import Session
from core import database
from core.database import get_db
from models.models import Document, DocumentChunk
from utilities.faiss_index import FaissIndex
from utilities.embeddings import embed_texts, embed_query
from core.role_based import admin_required, any_registered_user, admin_or_user
from core.secure import get_current_user
from fastapi.responses import JSONResponse
import os, uuid, traceback

# try to use langchain_text_splitters if available for better splitting
try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    _HAS_SPLITTER = True
except Exception:
    _HAS_SPLITTER = False

router = APIRouter()
DIM = 1536  # set to embedding dim (text-embedding-3-small is 1536)
faiss_index = FaissIndex(dim=DIM)
get_db = database.get_db

def chunk_text(text: str, chunk_size: int = 800, overlap: int = 200):
    """
    Returns list[str] chunks. Prefer LangChain splitter if installed.
    """
    if _HAS_SPLITTER:
        splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=overlap)
        return splitter.split_text(text)
    # fallback naive whitespace splitter
    tokens = text.split()
    chunks = []
    i = 0
    while i < len(tokens):
        chunk = " ".join(tokens[i:i+chunk_size])
        chunks.append(chunk)
        i += chunk_size - overlap
    return chunks

@router.post("/upload", dependencies=[Depends(admin_required)])
def upload_document(file: UploadFile = File(...), db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    """
    Upload a document, extract text (pdf/docx/txt), split into chunks, embed and index via FAISS + save chunks to DB.
    """
    try:
        UPLOAD_DIR = "uploads/documents"
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        filename = f"{uuid.uuid4().hex}_{file.filename}"
        path = os.path.join(UPLOAD_DIR, filename)
        with open(path, "wb") as f:
            f.write(file.file.read())

        # extract text
        content = ""
        name = file.filename.lower()
        if name.endswith(".txt"):
            with open(path, "r", encoding="utf-8", errors="ignore") as fh:
                content = fh.read()
        elif name.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                pages = [p.extract_text() or "" for p in pdf.pages]
                content = "\n".join(pages)
        elif name.endswith(".docx"):
            import docx
            doc = docx.Document(path)
            content = "\n".join([p.text for p in doc.paragraphs])
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")

        # save Document row
        doc = Document(filename=filename, uploader_id=current_user.id)
        db.add(doc)
        db.commit()
        db.refresh(doc)

        # chunk text
        chunks = chunk_text(content)
        if not chunks:
            raise HTTPException(status_code=400, detail="No text extracted from document")

        # create embeddings
        embeddings = embed_texts(chunks)

        # persist chunks and add to faiss
        metadata_ids = []
        for idx, (txt, emb) in enumerate(zip(chunks, embeddings)):
            chunk_row = DocumentChunk(document_id=doc.id, chunk_index=idx, text=txt, embedding=emb)
            db.add(chunk_row)
            db.commit()
            db.refresh(chunk_row)
            metadata_ids.append(chunk_row.id)

        faiss_index.add(embeddings, metadata_ids)
        faiss_index.save()

        return JSONResponse(
            status_code=status.HTTP_200_OK,
            content={
            "message": "Uploaded and indexed",
            "document_id": doc.id,
            "filename": filename,
            "chunks": len(chunks)
        })

    except Exception as e:
        # rollback DB session and return understandable error
        try:
            db.rollback()
        except Exception:
            pass
        tb = traceback.format_exc()
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Unable to upload document. {str(e)}\n\n{tb}")

@router.get("/search", dependencies=[Depends(admin_or_user)])
def search(q: str, k: int = 5, db: Session = Depends(get_db), current_user: dict = Depends(get_current_user)):
    if not q:
        raise HTTPException(status_code=400, detail="Query is required")
    try:
        q_emb = embed_query(q)
        results = faiss_index.search(q_emb, k=k)
        ids = [r["metadata_id"] for r in results]
        chunks = db.query(DocumentChunk).filter(DocumentChunk.id.in_(ids)).all()
        chunk_map = {c.id: c for c in chunks}
        out = []
        for r in results:
            c = chunk_map.get(r["metadata_id"])
            if not c:
                continue
            out.append({
                "chunk_id": c.id,
                "document_id": c.document_id,
                "text_snippet": (c.text[:500] + "...") if len(c.text) > 500 else c.text,
                "score": r["score"]
            })
        return {"query": q, "results": out}
    except Exception as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=f"Search failed. {str(e)}")
